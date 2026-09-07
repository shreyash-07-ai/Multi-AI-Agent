from pathlib import Path
from docx import Document
from pptx import Presentation

def validate(artifacts):
    report = {"status": "passed", "checks": []}
    for kind, path in artifacts.items():
        p = Path(path)
        ok = p.exists() and p.stat().st_size > 0
        check = {"artifact": kind, "exists": ok, "path": str(p)}
        if ok and p.suffix.lower() == ".docx":
            doc = Document(p)
            check["paragraphs"] = len(doc.paragraphs)
            check["editable"] = True
        if ok and p.suffix.lower() == ".pptx":
            prs = Presentation(p)
            check["slides"] = len(prs.slides)
            check["editable"] = True
        report["checks"].append(check)
        if not ok:
            report["status"] = "failed"
    return report
