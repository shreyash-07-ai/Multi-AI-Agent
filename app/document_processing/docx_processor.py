from docx import Document

def analyze_docx(path):
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    styles = {}
    for p in doc.paragraphs:
        if p.text.strip():
            styles[p.style.name] = styles.get(p.style.name, 0) + 1
    tables = []
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([c.text.strip() for c in row.cells])
        tables.append(rows)
    return {
        "type": "docx",
        "paragraphs": paragraphs,
        "styles": styles,
        "tables": tables,
        "paragraph_count": len(paragraphs),
        "table_count": len(tables),
    }

def extract_text(path):
    info = analyze_docx(path)
    return "\n".join(info["paragraphs"])
