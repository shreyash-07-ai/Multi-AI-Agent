from pptx import Presentation

def analyze_pptx(path):
    prs = Presentation(path)
    slides = []
    for idx, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        slides.append({"slide": idx, "texts": texts, "layout": slide.slide_layout.name})
    return {
        "type": "pptx",
        "slide_count": len(slides),
        "slides": slides,
        "slide_width": prs.slide_width,
        "slide_height": prs.slide_height,
        "layout_names": [l.name for l in prs.slide_layouts],
    }

def extract_text(path):
    info = analyze_pptx(path)
    chunks = []
    for s in info["slides"]:
        chunks.append(f"Slide {s['slide']}: " + " | ".join(s["texts"]))
    return "\n".join(chunks)
