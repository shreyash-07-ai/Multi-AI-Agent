from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _clear_document_body(doc):
    body = doc._element.body

    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def _set_cell_shading(cell, fill="D9EAF7"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.25 * level)
    paragraph.paragraph_format.space_after = Pt(4)

    run = paragraph.add_run(str(text))
    run.font.size = Pt(10.5)

    return paragraph


def _add_section(doc, section):
    heading = section.get("heading", "Section")

    p = doc.add_paragraph()
    p.style = "Heading 1"
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)

    run = p.add_run(str(heading))
    run.bold = True
    run.font.size = Pt(15)

    bullets = section.get("bullets", [])

    if isinstance(bullets, str):
        bullets = [bullets]

    for bullet in bullets:
        if isinstance(bullet, dict):
            text = bullet.get("text", bullet.get("content", ""))
        else:
            text = bullet

        if text:
            _add_bullet(doc, text)


def _add_citations(doc, citations):
    if not citations:
        return

    p = doc.add_paragraph()
    p.style = "Heading 1"

    run = p.add_run("Sources & Traceability")
    run.bold = True
    run.font.size = Pt(15)

    for source in citations:
        if isinstance(source, dict):
            title = source.get("title", "Source")
            url = source.get("url", "")

            text = title
            if url:
                text += f" — {url}"
        else:
            text = str(source)

        _add_bullet(doc, text)


def generate_docx(content, output_path, template_path=None):
    """
    Generate a properly formatted DOCX from structured Gemini JSON.

    Expected content:
    {
        "title": "...",
        "executive_summary": "...",
        "sections": [
            {
                "heading": "...",
                "bullets": ["...", "..."]
            }
        ],
        "citations": [...]
    }
    """

    if not isinstance(content, dict):
        content = {
            "title": "AI-Generated Business Proposal",
            "executive_summary": str(content),
            "sections": [],
            "citations": [],
        }

    if template_path and Path(template_path).exists():
        doc = Document(template_path)
        _clear_document_body(doc)
    else:
        doc = Document()

    # ---------------------------------------------------------
    # TITLE
    # ---------------------------------------------------------

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)

    run = title.add_run(
        str(content.get("title", "AI-Generated Business Proposal"))
    )
    run.bold = True
    run.font.size = Pt(24)

    # ---------------------------------------------------------
    # EXECUTIVE SUMMARY
    # ---------------------------------------------------------

    heading = doc.add_paragraph()
    heading.style = "Heading 1"

    run = heading.add_run("Executive Summary")
    run.bold = True
    run.font.size = Pt(15)

    summary = content.get("executive_summary", "")

    if summary:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)

        run = p.add_run(str(summary))
        run.font.size = Pt(10.5)

    # ---------------------------------------------------------
    # SECTIONS
    # ---------------------------------------------------------

    sections = content.get("sections", [])

    if isinstance(sections, dict):
        sections = [sections]

    for section in sections:
        if isinstance(section, dict):
            _add_section(doc, section)

    # ---------------------------------------------------------
    # OPTIONAL KEY-VALUE / METADATA
    # ---------------------------------------------------------

    metadata = content.get("metadata")

    if isinstance(metadata, dict) and metadata:
        heading = doc.add_paragraph()
        heading.style = "Heading 1"

        run = heading.add_run("Key Information")
        run.bold = True
        run.font.size = Pt(15)

        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"

        table.rows[0].cells[0].text = "Item"
        table.rows[0].cells[1].text = "Value"

        _set_cell_shading(table.rows[0].cells[0])
        _set_cell_shading(table.rows[0].cells[1])

        for key, value in metadata.items():
            cells = table.add_row().cells
            cells[0].text = str(key)
            cells[1].text = str(value)

    # ---------------------------------------------------------
    # SOURCES
    # ---------------------------------------------------------

    _add_citations(doc, content.get("citations", []))

    # ---------------------------------------------------------
    # SAVE
    # ---------------------------------------------------------

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc.save(str(output_path))

    return str(output_path)