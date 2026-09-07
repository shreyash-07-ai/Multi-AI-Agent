from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


# ============================================================
# PRESENTATION SETTINGS
# ============================================================
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

LEFT = Inches(0.72)
RIGHT = Inches(0.72)
CONTENT_WIDTH = Inches(11.893)

TITLE_TOP = Inches(0.42)
TITLE_HEIGHT = Inches(0.62)
CONTENT_TOP = Inches(1.35)
CONTENT_HEIGHT = Inches(5.55)
FOOTER_Y = Inches(7.08)

MAX_BULLETS_PER_SLIDE = 4

TEXT_LEFT = Inches(1.02)
TEXT_TOP = Inches(1.68)
TEXT_WIDTH = Inches(11.25)
TEXT_HEIGHT = Inches(4.92)

BG_COLOR = RGBColor(18, 22, 30)
CARD_COLOR = RGBColor(29, 35, 46)
LINE_COLOR = RGBColor(62, 72, 90)
TITLE_COLOR = RGBColor(245, 247, 250)
ACCENT_COLOR = RGBColor(88, 166, 255)
MUTED_COLOR = RGBColor(175, 184, 198)


# ============================================================
# HELPERS
# ============================================================
def _delete_all_slides(prs):
    for slide_id in list(prs.slides._sldIdLst):
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)


def _get_blank_layout(prs):
    for layout in prs.slide_layouts:
        if layout.name.lower() == "blank":
            return layout
    return prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]


def _set_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR


def _clean_text(text):
    if text is None:
        return ""
    return (
        str(text)
        .replace("```json", "")
        .replace("```", "")
        .replace("**", "")
        .replace("__", "")
        .strip()
    )


def _normalize_bullets(bullets):
    if bullets is None:
        return []
    if isinstance(bullets, str):
        bullets = [bullets]
    if not isinstance(bullets, list):
        return []

    result = []
    for item in bullets:
        if isinstance(item, dict):
            item = item.get("text", item.get("content", ""))
        item = _clean_text(item)
        if item:
            result.append(item)
    return result


def _fit_font_size(bullets):
    total_chars = sum(len(x) for x in bullets)
    longest = max((len(x) for x in bullets), default=0)

    if len(bullets) <= 2 and total_chars < 420 and longest < 230:
        return 21
    if len(bullets) <= 3 and total_chars < 620 and longest < 300:
        return 19
    if total_chars < 850 and longest < 390:
        return 17
    return 15


def _add_text_box(
    slide, text, x, y, w, h, size=18, bold=False,
    color=TITLE_COLOR, align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = valign

    p = tf.paragraphs[0]
    p.text = _clean_text(text)
    p.alignment = align
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def _add_footer(slide, number, total):
    _add_text_box(
        slide, f"{number} / {total}", Inches(11.45), FOOTER_Y,
        Inches(1.1), Inches(0.22), size=9,
        color=MUTED_COLOR, align=PP_ALIGN.RIGHT,
    )


def _add_header(slide, title):
    _add_text_box(
        slide, title, LEFT, TITLE_TOP, CONTENT_WIDTH, TITLE_HEIGHT,
        size=27, bold=True, color=TITLE_COLOR, align=PP_ALIGN.LEFT,
    )

    accent = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, LEFT, Inches(1.10), Inches(0.72), Inches(0.045)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT_COLOR
    accent.line.fill.background()


def _add_content_card(slide, top=CONTENT_TOP, height=CONTENT_HEIGHT):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        LEFT, top, CONTENT_WIDTH, height,
    )
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_COLOR
    card.line.color.rgb = LINE_COLOR
    card.line.width = Pt(1)
    return card


def _add_bullet_list(slide, bullets):
    _add_content_card(slide)

    text_box = slide.shapes.add_textbox(
        TEXT_LEFT, TEXT_TOP, TEXT_WIDTH, TEXT_HEIGHT
    )
    tf = text_box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP

    font_size = _fit_font_size(bullets)
    space_after = 13 if len(bullets) <= 3 else 9

    for index, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = f"• {bullet}"
        p.font.name = "Aptos"
        p.font.size = Pt(font_size)
        p.font.color.rgb = TITLE_COLOR
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        p.line_spacing = 1.08
        p.level = 0

    return text_box


def _add_summary_card(slide, summary):
    _add_content_card(slide, top=Inches(1.55), height=Inches(4.95))
    summary = _clean_text(summary)

    if len(summary) < 600:
        size = 20
    elif len(summary) < 900:
        size = 17
    else:
        size = 15

    _add_text_box(
        slide, summary,
        Inches(1.15), Inches(1.95), Inches(11.05), Inches(4.15),
        size=size, color=TITLE_COLOR,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
    )


# ============================================================
# TITLE SLIDE
# ============================================================
def _create_title_slide(prs, content, layout, number, total):
    """
    Dedicated title slide.

    IMPORTANT: the executive summary is intentionally NOT placed here.
    A title slide should visually read as a title/cover page rather than
    as a large content card behind the title.
    """
    slide = prs.slides.add_slide(layout)
    _set_background(slide)

    title = _clean_text(
        content.get("title", "AI-Generated Business Proposal")
    )
    company = _clean_text(
        content.get("company_name", content.get("company", ""))
    )
    project = _clean_text(
        content.get("project_name", content.get("topic", ""))
    )

    # Small presentation label.
    _add_text_box(
        slide, "AI-GENERATED BUSINESS PROPOSAL",
        Inches(1.0), Inches(1.28), Inches(11.33), Inches(0.35),
        size=11, bold=True, color=ACCENT_COLOR,
        align=PP_ALIGN.CENTER,
    )

    # Main title — centered with generous whitespace.
    title_size = 34 if len(title) < 55 else 29 if len(title) < 85 else 25
    title_height = Inches(1.25)
    _add_text_box(
        slide, title,
        Inches(1.0), Inches(1.85), Inches(11.33), title_height,
        size=title_size, bold=True, color=TITLE_COLOR,
        align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE,
    )

    # Accent divider.
    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(5.85), Inches(3.22), Inches(1.63), Inches(0.055),
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = ACCENT_COLOR
    divider.line.fill.background()

    # Topic/project subtitle.
    if project:
        _add_text_box(
            slide, project,
            Inches(1.35), Inches(3.58), Inches(10.63), Inches(0.58),
            size=19, color=WHITE, align=PP_ALIGN.CENTER,
        )

    # Company appears below the topic, not inside a large content box.
    if company:
        _add_text_box(
            slide, company,
            Inches(1.35), Inches(4.28), Inches(10.63), Inches(0.45),
            size=14, color=MUTED_COLOR, align=PP_ALIGN.CENTER,
        )

    # Subtle bottom line makes the cover feel intentional without adding
    # another content block.
    bottom = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(2.2), Inches(6.35), Inches(8.93), Inches(0.018),
    )
    bottom.fill.solid()
    bottom.fill.fore_color.rgb = LINE_COLOR
    bottom.line.fill.background()

    _add_footer(slide, number, total)
    return slide


# ============================================================
# CONTENT SLIDES
# ============================================================
def _create_content_slide(prs, title, bullets, layout, number, total):
    slide = prs.slides.add_slide(layout)
    _set_background(slide)
    _add_header(slide, title)
    _add_bullet_list(slide, bullets)
    _add_footer(slide, number, total)
    return slide


def _create_summary_slide(prs, summary, layout, number, total):
    slide = prs.slides.add_slide(layout)
    _set_background(slide)
    _add_header(slide, "Executive Summary")
    _add_summary_card(slide, summary)
    _add_footer(slide, number, total)
    return slide


# ============================================================
# SLIDE DATA
# ============================================================
def _slides_from_gemini(content):
    """Prefer Gemini's explicit slide plan when it exists."""
    generated = content.get("slides", [])
    result = []

    if not isinstance(generated, list):
        return result

    for item in generated:
        if not isinstance(item, dict):
            continue

        title = _clean_text(item.get("title", "Section"))
        bullets = _normalize_bullets(item.get("bullets", []))
        if not bullets:
            continue

        for i in range(0, len(bullets), MAX_BULLETS_PER_SLIDE):
            chunk = bullets[i:i + MAX_BULLETS_PER_SLIDE]
            if not chunk:
                continue

            if len(bullets) > MAX_BULLETS_PER_SLIDE:
                part = i // MAX_BULLETS_PER_SLIDE + 1
                parts = (
                    len(bullets) + MAX_BULLETS_PER_SLIDE - 1
                ) // MAX_BULLETS_PER_SLIDE
                slide_title = f"{title} ({part}/{parts})"
            else:
                slide_title = title

            result.append({"title": slide_title, "bullets": chunk})

    return result


def _slides_from_sections(content):
    result = []
    sections = content.get("sections", [])

    if not isinstance(sections, list):
        return result

    for section in sections:
        if not isinstance(section, dict):
            continue

        heading = _clean_text(
            section.get("heading", section.get("title", "Section"))
        )
        bullets = _normalize_bullets(section.get("bullets", []))

        for i in range(0, len(bullets), MAX_BULLETS_PER_SLIDE):
            chunk = bullets[i:i + MAX_BULLETS_PER_SLIDE]
            if not chunk:
                continue

            if len(bullets) > MAX_BULLETS_PER_SLIDE:
                part = i // MAX_BULLETS_PER_SLIDE + 1
                parts = (
                    len(bullets) + MAX_BULLETS_PER_SLIDE - 1
                ) // MAX_BULLETS_PER_SLIDE
                title = f"{heading} ({part}/{parts})"
            else:
                title = heading

            result.append({"title": title, "bullets": chunk})

    return result


def _slides_from_citations(content):
    citations = content.get("citations", [])
    if not isinstance(citations, list):
        return []

    sources = []
    for citation in citations:
        if isinstance(citation, dict):
            title = _clean_text(citation.get("title", "Source"))
            url = _clean_text(citation.get("url", ""))
            sources.append(f"{title} — {url}" if url else title)
        else:
            text = _clean_text(citation)
            if text:
                sources.append(text)

    result = []
    for i in range(0, len(sources), MAX_BULLETS_PER_SLIDE):
        chunk = sources[i:i + MAX_BULLETS_PER_SLIDE]
        if chunk:
            result.append({"title": "Sources & References", "bullets": chunk})
    return result


# ============================================================
# GENERATE PPTX
# ============================================================
def generate_pptx(content, output_path, template_path=None, target_slides=12):
    if template_path and Path(template_path).exists():
        prs = Presentation(template_path)
    else:
        prs = Presentation()

    # Predictable 16:9 canvas for consistent alignment.
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    _delete_all_slides(prs)
    blank_layout = _get_blank_layout(prs)

    # Gemini's explicit slides are preferred. Sections are the fallback.
    slide_data = _slides_from_gemini(content)
    if not slide_data:
        slide_data = _slides_from_sections(content)

    if not slide_data:
        slide_data = [{
            "title": "Overview",
            "bullets": [_clean_text(content.get("executive_summary", "")) or "No additional information available."],
        }]

    citation_slides = _slides_from_citations(content)
    if citation_slides:
        slide_data.extend(citation_slides)

    # The cover is always slide 1. Executive summary is a separate slide 2.
    summary = _clean_text(content.get("executive_summary", ""))
    has_summary = bool(summary)
    total = 1 + (1 if has_summary else 0) + len(slide_data)

    _create_title_slide(prs, content, blank_layout, 1, total)

    next_number = 2
    if has_summary:
        _create_summary_slide(
            prs, summary, blank_layout, next_number, total
        )
        next_number += 1

    for item in slide_data:
        _create_content_slide(
            prs,
            item["title"],
            item["bullets"],
            blank_layout,
            next_number,
            total,
        )
        next_number += 1

    prs.save(output_path)
    return str(output_path)
